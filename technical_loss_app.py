"""
Technical Loss Automation Simulator - Streamlit App
=====================================================
Feeder Consumption + Installation Master + Billing Cycle Report
-> Feeder-wise loss report, with an expandable consumer-detail Excel export.

Run with:  streamlit run technical_loss_app.py
For the 1.6GB installation file, .streamlit/config.toml (shipped alongside
this file) raises Streamlit's upload cap to 2000MB.

LOSS FORMULA
------------
Loss (kWh) = Total Feeder Consumption - Total Consumer Consumption
Loss (%)   = Loss (kWh) / Total Feeder Consumption x 100

KNOWN DATA-QUALITY FLAGS (surfaced in the UI and Excel export)
----------------------------------------------------------------
1. Unmatched consumers - billed but not found in the installation master
   for this feeder. Their consumption is excluded from the feeder sum,
   which inflates that feeder's apparent loss.
2. Missing previous-month reading - delta can't be computed, so that
   consumer's consumption is NaN and is dropped from the sum (same
   inflate-loss effect as #1).
3. Negative monthly delta - current reading < previous reading, usually
   a meter swap/rollover mid-month. Reduces the consumer sum artificially.
"""

import io
import json
import requests
import pandas as pd
import streamlit as st
from openpyxl.utils.dataframe import dataframe_to_rows

st.set_page_config(page_title="Technical Loss Automation Simulator", layout="wide", page_icon="⚡")

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
INSTALL_COLS_NEEDED = ["Feeder Code", "Feeder Name", "Consumer Number (KNO)", "DT Code"]
INSTALL_CHUNK_SIZE = 50_000

BILLING_COLS_NEEDED = [
    "feeder_name", "consumer_name", "consumer_number",
    "Active_energy_import", "Active_energy_export",
]

# ---------------------------------------------------------------------------
# Simple blue & white theme
# ---------------------------------------------------------------------------
st.markdown("""
<style>
:root {
    --blue: #1E5FCB;
    --blue-dark: #14428F;
    --blue-light: #EAF1FC;
    --border: #D8E2F0;
    --text: #1B2430;
}
.stApp { background-color: #FFFFFF; }
h1, h2, h3, h4 { color: var(--blue-dark) !important; }
p, span, label { color: var(--text); }
section[data-testid="stSidebar"] { background-color: var(--blue-light); border-right: 1px solid var(--border); }
[data-testid="stMetricValue"] { color: var(--blue-dark) !important; }
.stButton>button, .stDownloadButton>button {
    background-color: var(--blue) !important; color: #FFFFFF !important;
    border: none !important; border-radius: 6px !important; font-weight: 600 !important;
}
.stButton>button:hover, .stDownloadButton>button:hover { background-color: var(--blue-dark) !important; }
[data-testid="stFileUploaderDropzone"], [data-testid="stFileUploader"] section {
    background-color: var(--blue-light) !important; border: 1.5px dashed var(--blue) !important; border-radius: 8px !important;
}
[data-testid="stAlert"] { border-radius: 8px !important; }
[data-testid="stDataFrame"] { border: 1px solid var(--border) !important; border-radius: 8px; }
</style>
""", unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------
def parse_feeder_consumption_json(data: dict) -> pd.DataFrame:
    """Shared parser for the {'MonthlyEnergy': [...]} shape, used by both
    the file upload path and the API fetch path so the logic stays in one place."""
    df = pd.DataFrame(data["MonthlyEnergy"])
    df["unin"] = df["unin"].astype(str).str.strip()
    df["import_kwh"] = pd.to_numeric(df["import_kwh"], errors="coerce")
    df["feedername"] = df["feedername"].astype(str).str.strip()
    return df[["unin", "feedername", "import_kwh"]].drop_duplicates(subset="unin")


def load_feeder_consumption(uploaded_file) -> pd.DataFrame:
    name = uploaded_file.name.lower()
    if name.endswith(".json"):
        data = json.load(uploaded_file)
        return parse_feeder_consumption_json(data)
    try:
        df = pd.read_csv(uploaded_file, encoding="utf-8-sig")
    except UnicodeDecodeError:
        uploaded_file.seek(0)
        df = pd.read_csv(uploaded_file, encoding="latin-1")  # fallback for non-UTF-8 exports
    df["unin"] = df["unin"].astype(str).str.strip()
    df["import_kwh"] = pd.to_numeric(df["import_kwh"], errors="coerce")
    df["feedername"] = df["feedername"].astype(str).str.strip()
    return df[["unin", "feedername", "import_kwh"]].drop_duplicates(subset="unin")


def fetch_feeder_consumption_from_api(url_template: str, year: int, month: int, subdivision: str, timeout: int = 30) -> pd.DataFrame:
    """
    Calls the feeder consumption API for a given year + month + sub-division and
    parses the response with the same logic used for uploaded JSON files.
    url_template must contain {year}, {month} and {subdivision} placeholders, e.g.:
        http://10.1.1.1/api/feeder-consumption?year={year}&month={month}&subdivision={subdivision}
    """
    url = url_template.format(year=year, month=month, subdivision=subdivision)
    resp = requests.get(url, timeout=timeout)
    resp.raise_for_status()
    data = resp.json()
    return parse_feeder_consumption_json(data)


def load_billing_report(uploaded_file) -> pd.DataFrame:
    raw = uploaded_file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")  # fallback for non-UTF-8 exports
    df = pd.read_csv(io.StringIO(text), skiprows=4)
    missing = [c for c in BILLING_COLS_NEEDED if c not in df.columns]
    if missing:
        raise ValueError(f"Billing file '{uploaded_file.name}' is missing columns: {missing}")
    df = df[BILLING_COLS_NEEDED].copy()
    df["consumer_number"] = df["consumer_number"].astype(str).str.strip()
    df["Active_energy_import"] = pd.to_numeric(df["Active_energy_import"], errors="coerce")
    df["Active_energy_export"] = pd.to_numeric(df["Active_energy_export"], errors="coerce")
    df["net_energy"] = df["Active_energy_import"] - df["Active_energy_export"].fillna(0)
    return df


def load_installation_chunked(uploaded_file, chunk_size=INSTALL_CHUNK_SIZE) -> pd.DataFrame:
    chunks = []
    progress = st.progress(0.0, text="Reading installation file...")
    # latin-1 (ISO-8859-1) instead of utf-8-sig: it can decode any byte value without
    # raising, unlike utf-8-sig which crashes mid-file on non-UTF-8 exports (common from
    # Excel/Windows systems). Safe here since we only keep numeric/ASCII columns.
    reader = pd.read_csv(
        uploaded_file, usecols=lambda c: c.strip() in INSTALL_COLS_NEEDED,
        dtype=str, chunksize=chunk_size, encoding="latin-1", on_bad_lines="skip",
    )
    seen = 0
    for i, chunk in enumerate(reader):
        chunk.columns = [c.strip() for c in chunk.columns]
        chunk = chunk.dropna(subset=["Consumer Number (KNO)", "Feeder Code"])
        chunk["Consumer Number (KNO)"] = chunk["Consumer Number (KNO)"].str.strip()
        chunk["Feeder Code"] = chunk["Feeder Code"].str.strip()
        chunks.append(chunk.drop_duplicates(subset="Consumer Number (KNO)"))
        seen += len(chunk)
        progress.progress(min(0.95, (i + 1) * 0.03), text=f"Processed ~{seen:,} rows...")
    progress.progress(1.0, text="Installation file loaded.")
    return pd.concat(chunks, ignore_index=True).drop_duplicates(subset="Consumer Number (KNO)")


# ---------------------------------------------------------------------------
# Core calculation - returns feeder-wise report AND consumer-level detail
# (detail is needed for the expandable Excel export + per-feeder stats)
# ---------------------------------------------------------------------------
def compute_report(feeder_df, install_df, current_bill_df, previous_bill_df):
    merged = current_bill_df.merge(
        previous_bill_df[["consumer_number", "Active_energy_import", "Active_energy_export", "net_energy"]]
        .rename(columns={"Active_energy_import": "import_prev", "Active_energy_export": "export_prev",
                          "net_energy": "net_energy_prev"}),
        on="consumer_number", how="left",
    )
    merged["has_prev_reading"] = merged["net_energy_prev"].notna()
    merged["consumer_consumption"] = merged["net_energy"] - merged["net_energy_prev"]
    merged["negative_delta"] = merged["consumer_consumption"] < 0

    merged = merged.merge(
        install_df[["Feeder Code", "Consumer Number (KNO)"]],
        left_on="consumer_number", right_on="Consumer Number (KNO)", how="left",
    )
    merged["matched_to_feeder"] = merged["Feeder Code"].notna()

    consumer_detail = merged.rename(columns={
        "consumer_name": "Consumer Name", "consumer_number": "Consumer Number",
        "feeder_name": "Feeder Name (billing)", "Feeder Code": "Feeder Code",
        "Active_energy_import": "Import (kWh)", "Active_energy_export": "Export (kWh)",
        "net_energy": "Actual Energy (kWh)", "consumer_consumption": "Monthly Consumption (kWh)",
    })[["Consumer Name", "Consumer Number", "Feeder Name (billing)", "Feeder Code",
        "Import (kWh)", "Export (kWh)", "Actual Energy (kWh)", "Monthly Consumption (kWh)",
        "matched_to_feeder", "has_prev_reading", "negative_delta"]]

    feeder_agg = (
        merged.dropna(subset=["Feeder Code"])
        .groupby("Feeder Code")
        .agg(total_consumer_consumption=("consumer_consumption", "sum"),
             consumer_count=("consumer_number", "count"))
        .reset_index()
    )

    # Installed-but-never-billed: consumers present in the installation master
    # for a feeder but absent from this month's billing report entirely.
    billed_ids = set(current_bill_df["consumer_number"])
    never_billed = install_df[~install_df["Consumer Number (KNO)"].isin(billed_ids)].rename(columns={
        "Feeder Code": "Feeder Code", "Feeder Name": "Feeder Name (installation)",
        "Consumer Number (KNO)": "Consumer Number",
    })[["Consumer Number", "Feeder Code", "Feeder Name (installation)"]]

    report = feeder_df.merge(feeder_agg, left_on="unin", right_on="Feeder Code", how="left")
    report = report.drop(columns=["Feeder Code"])
    report["total_consumer_consumption"] = report["total_consumer_consumption"].fillna(0)
    report["consumer_count"] = report["consumer_count"].fillna(0).astype(int)
    report["loss_kwh"] = report["import_kwh"] - report["total_consumer_consumption"]
    report["loss_pct"] = (report["loss_kwh"] / report["import_kwh"]) * 100

    report = report.rename(columns={
        "feedername": "Feeder Name", "unin": "Feeder Code",
        "import_kwh": "Total Feeder Consumption (kWh)",
        "total_consumer_consumption": "Total Consumer Consumption (kWh)",
        "consumer_count": "Number of Consumers",
        "loss_kwh": "Loss (kWh)", "loss_pct": "Loss (%)",
    })[["Feeder Name", "Feeder Code", "Number of Consumers", "Total Feeder Consumption (kWh)",
        "Total Consumer Consumption (kWh)", "Loss (kWh)", "Loss (%)"]]

    return report, consumer_detail, never_billed


# ---------------------------------------------------------------------------
# Excel export: Feeder-wise summary sheet + an expandable
# "Feeder + Consumer Detail" sheet using Excel's native row grouping.
# Click the "+" next to a feeder row to expand its consumers below it.
# ---------------------------------------------------------------------------
def to_excel_bytes(report_df: pd.DataFrame, consumer_df: pd.DataFrame, never_billed_df: pd.DataFrame) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()

    # --- Sheet 1: feeder-wise summary, with a Exporting Consumers section appended below ---
    ws1 = wb.active
    ws1.title = "Feeder-wise Report"
    for r in dataframe_to_rows(report_df, index=False, header=True):
        ws1.append(r)
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 20

    negative_df = consumer_df[consumer_df["negative_delta"]][
        ["Consumer Name", "Consumer Number", "Feeder Name (billing)", "Feeder Code",
         "Monthly Consumption (kWh)"]
    ]
    ws1.append([])
    ws1.append(["Exporting Consumers (current reading below previous month — check for meter swap/rollover)"])
    section_title_row = ws1.max_row
    ws1.cell(row=section_title_row, column=1).font = ws1.cell(row=section_title_row, column=1).font.copy(bold=True)
    if negative_df.empty:
        ws1.append(["No consumers with a negative monthly delta this month."])
    else:
        for r in dataframe_to_rows(negative_df, index=False, header=True):
            ws1.append(r)
        neg_header_row = section_title_row + 1
        for cell in ws1[neg_header_row]:
            cell.font = cell.font.copy(bold=True)

    # --- Sheet 2: expandable feeder -> consumer detail ---
    ws2 = wb.create_sheet("Feeder + Consumer Detail")
    ws2.sheet_properties.outlinePr.summaryBelow = False  # group control sits above the detail rows
    detail_headers = ["Consumer Name", "Consumer Number", "Import (kWh)", "Export (kWh)",
                       "Actual Energy (kWh)", "Monthly Consumption (kWh)"]
    ws2.append(["Feeder Name", "Feeder Code", "Total Feeder Consumption (kWh)",
                "Total Consumer Consumption (kWh)", "Loss (kWh)", "Loss (%)"] + [""] * 2)
    header_row_idx = ws2.max_row
    for cell in ws2[header_row_idx]:
        cell.font = cell.font.copy(bold=True)

    consumer_by_feeder = {code: grp for code, grp in consumer_df.groupby("Feeder Code")}

    for _, frow in report_df.sort_values("Loss (%)", ascending=False).iterrows():
        ws2.append([
            frow["Feeder Name"], frow["Feeder Code"], frow["Total Feeder Consumption (kWh)"],
            frow["Total Consumer Consumption (kWh)"], frow["Loss (kWh)"], frow["Loss (%)"],
        ])
        feeder_row_cells = ws2[ws2.max_row]
        for cell in feeder_row_cells:
            cell.font = cell.font.copy(bold=True)

        sub = consumer_by_feeder.get(str(frow["Feeder Code"]))
        if sub is None or sub.empty:
            continue

        # Small sub-header for the detail columns, still inside the collapsible group
        ws2.append(detail_headers)
        sub_header_row = ws2.max_row
        ws2.row_dimensions[sub_header_row].outlineLevel = 1
        ws2.row_dimensions[sub_header_row].hidden = True

        for _, crow in sub.iterrows():
            ws2.append([
                crow["Consumer Name"], crow["Consumer Number"], crow["Import (kWh)"],
                crow["Export (kWh)"], crow["Actual Energy (kWh)"], crow["Monthly Consumption (kWh)"],
            ])
            ws2.row_dimensions[ws2.max_row].outlineLevel = 1
            ws2.row_dimensions[ws2.max_row].hidden = True

    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 20

    # --- Sheets 3-6: each data-quality flag gets its own section, so
    # you can audit "why is this feeder's loss high" one cause at a time ---
    def write_sheet(name, df):
        ws = wb.create_sheet(name)
        if df.empty:
            ws.append(["No rows for this flag."])
            return
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24

    unmatched_df = consumer_df[~consumer_df["matched_to_feeder"]][
        ["Consumer Name", "Consumer Number", "Feeder Name (billing)"]
    ]
    no_prev_df = consumer_df[~consumer_df["has_prev_reading"]][
        ["Consumer Name", "Consumer Number", "Feeder Name (billing)", "Feeder Code"]
    ]
    negative_df = consumer_df[consumer_df["negative_delta"]][
        ["Consumer Name", "Consumer Number", "Feeder Name (billing)", "Feeder Code",
         "Monthly Consumption (kWh)"]
    ]

    write_sheet("Flag - Unmatched to Feeder", unmatched_df)
    write_sheet("Flag - No Previous Reading", no_prev_df)
    write_sheet("Flag - Exporting Consumers", negative_df)
    write_sheet("Flag - Installed Never Billed", never_billed_df)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
st.markdown("## ⚡ Technical Loss Automation Simulator")
st.caption("Feeder-wise loss = Total Feeder Consumption − Total Consumer Consumption")
st.divider()

with st.sidebar:
    st.markdown("### 1. Installation Master (large file)")
    install_file = st.file_uploader("WC Consumer Installation file (CSV, up to ~1.6GB)", type=["csv"], key="install")

    st.markdown("### 2. Feeder Consumption")
    feeder_source = st.radio("Source", ["Upload File(s)", "Fetch from API"], key="feeder_source")

    if feeder_source == "Upload File(s)":
        feeder_files = st.file_uploader("One file per month (CSV or JSON)", type=["csv", "json"],
                                         accept_multiple_files=True, key="feeder")
    else:
        feeder_files = None
        st.caption("URL must contain `{year}`, `{month}` and `{subdivision}` placeholders — see help text below the fetch button.")
        api_url_template = st.text_input(
            "Feeder Consumption API URL",
            value=st.session_state.get(
                "api_url_template",
                "https://jdvvnl.myxenius.com/feeder/V1/API/gtMnthlyCnsmptn/{month}/{year}/{subdivision}",
            ),
            key="api_url_template",
        )
        subdivision_code = st.text_input("Sub-division Code", key="subdivision_code")
        year_to_fetch = st.number_input("Year", min_value=2000, max_value=2100, value=2026, step=1, key="year_to_fetch")
        months_to_fetch = st.multiselect(
            "Month(s) to fetch (1 = Jan ... 12 = Dec)", list(range(1, 13)), key="months_to_fetch"
        )
        fetch_clicked = st.button("Fetch from API")
        with st.expander("How to paste the API URL"):
            st.markdown(
                "Paste your endpoint exactly as you'd call it, but replace the actual "
                "year/month/sub-division values with `{year}`, `{month}`, `{subdivision}` — "
                "works whether they're query params (`?year=...`) or path segments "
                "(`/gtMnthlyCnsmptn/{month}/{year}/{subdivision}`), as in the default above.\n\n"
                "The app expects a **GET** request returning JSON shaped like your existing "
                "files: `{\"MonthlyEnergy\": [ {...}, {...} ]}`. If your API needs an API key "
                "or auth header, tell me the header name and I'll add a field — never paste "
                "the actual key/token in chat."
            )

    st.markdown("### 3. Billing Cycle Report")
    billing_files = st.file_uploader("One file per month (CSV)", type=["csv"],
                                      accept_multiple_files=True, key="billing")

# feeder_data maps a month label -> an already-loaded feeder consumption DataFrame,
# regardless of whether it came from an uploaded file or the API.
feeder_data = {}

if feeder_source == "Upload File(s)" and feeder_files:
    st.subheader("Label each feeder consumption file by month")
    cols = st.columns(len(feeder_files))
    for c, f in zip(cols, feeder_files):
        with c:
            label = st.text_input(f"Month for {f.name}", value=f.name, key=f"flabel_{f.name}")
        feeder_data[label] = load_feeder_consumption(f)

elif feeder_source == "Fetch from API":
    if fetch_clicked:
        if not subdivision_code:
            st.error("Enter a sub-division code before fetching.")
        elif not months_to_fetch:
            st.error("Select at least one month before fetching.")
        else:
            fetched = st.session_state.get("api_feeder_data", {})
            new_count = 0
            for m in sorted(months_to_fetch):
                try:
                    with st.spinner(f"Fetching {m}/{year_to_fetch}, sub-division {subdivision_code}..."):
                        df = fetch_feeder_consumption_from_api(api_url_template, year_to_fetch, m, subdivision_code)
                    fetched[f"{m}/{year_to_fetch} (subdivision {subdivision_code})"] = df
                    new_count += 1
                except requests.exceptions.RequestException as e:
                    st.error(f"Month {m}/{year_to_fetch}: API request failed — {e}")
                except (KeyError, ValueError) as e:
                    st.error(f"Month {m}/{year_to_fetch}: unexpected response format — {e}")
            if new_count:
                st.session_state["api_feeder_data"] = fetched
                st.success(f"Fetched {new_count} month(s). Total available: {len(fetched)}.")
    feeder_data = st.session_state.get("api_feeder_data", {})
    if feeder_data:
        st.caption("Fetched so far: " + ", ".join(feeder_data.keys()))

billing_data = {}
if billing_files:
    st.subheader("Label each billing report file by month")
    cols = st.columns(len(billing_files))
    for c, f in zip(cols, billing_files):
        with c:
            label = st.text_input(f"Month for {f.name}", value=f.name, key=f"blabel_{f.name}")
            billing_data[label] = f

st.divider()

if feeder_data and billing_data and install_file:
    col1, col2 = st.columns(2)
    with col1:
        current_month_label = st.selectbox("Current month (feeder consumption)", list(feeder_data.keys()))
        current_bill_label = st.selectbox("Current month (billing report)", list(billing_data.keys()))
    with col2:
        prev_bill_label = st.selectbox(
            "Previous month (billing report, baseline)", list(billing_data.keys()),
            index=min(1, len(billing_data) - 1) if len(billing_data) > 1 else 0,
        )

    if st.button("Run Loss Calculation"):
        feeder_df = feeder_data[current_month_label]
        with st.spinner("Loading billing reports..."):
            current_bill_df = load_billing_report(billing_data[current_bill_label])
            previous_bill_df = load_billing_report(billing_data[prev_bill_label])
        install_df = load_installation_chunked(install_file)
        with st.spinner("Computing loss..."):
            report, consumer_detail, never_billed = compute_report(feeder_df, install_df, current_bill_df, previous_bill_df)
        st.session_state["report"] = report
        st.session_state["consumer_detail"] = consumer_detail
        st.session_state["never_billed"] = never_billed

    if "report" in st.session_state:
        report = st.session_state["report"]
        consumer_detail = st.session_state["consumer_detail"]
        never_billed = st.session_state["never_billed"]

        total_feeder = report["Total Feeder Consumption (kWh)"].sum()
        total_consumer = report["Total Consumer Consumption (kWh)"].sum()
        total_loss = total_feeder - total_consumer
        overall_pct = (total_loss / total_feeder * 100) if total_feeder else 0

        unmatched = (~consumer_detail["matched_to_feeder"]).sum()
        no_prev = (~consumer_detail["has_prev_reading"]).sum()
        negative_delta = consumer_detail["negative_delta"].sum()
        never_billed_count = len(never_billed)

        c1, c2, c3 = st.columns(3)
        c1.metric("Total Feeder Consumption", f"{total_feeder:,.0f} kWh")
        c2.metric("Total Consumer Consumption", f"{total_consumer:,.0f} kWh")
        c3.metric("Total Loss", f"{total_loss:,.0f} kWh", f"{overall_pct:.2f}%")

        st.markdown("### ⚠️ Data Quality Flags (each is a separate section in the Excel export)")
        f1, f2, f3, f4 = st.columns(4)
        f1.metric("Unmatched to Feeder", f"{unmatched:,}")
        f2.metric("No Previous Reading", f"{no_prev:,}")
        f3.metric("Exporting Consumers", f"{negative_delta:,}")
        f4.metric("Installed, Never Billed", f"{never_billed_count:,}")

        st.markdown("### Feeder-wise Report")
        st.dataframe(
            report.sort_values("Loss (%)", ascending=False).style.format({
                "Total Feeder Consumption (kWh)": "{:,.0f}",
                "Total Consumer Consumption (kWh)": "{:,.0f}",
                "Loss (kWh)": "{:,.0f}", "Loss (%)": "{:.2f}",
            }),
            use_container_width=True, hide_index=True,
        )

        st.markdown("### Per-Feeder Stats")
        selected_feeder = st.selectbox("Select a feeder", report["Feeder Name"].tolist())
        frow = report[report["Feeder Name"] == selected_feeder].iloc[0]
        s1, s2, s3, s4 = st.columns(4)
        s1.metric("Number of Consumers", f"{int(frow['Number of Consumers']):,}")
        s2.metric("Feeder Consumption", f"{frow['Total Feeder Consumption (kWh)']:,.0f} kWh")
        s3.metric("Consumer Consumption", f"{frow['Total Consumer Consumption (kWh)']:,.0f} kWh")
        s4.metric("Loss", f"{frow['Loss (kWh)']:,.0f} kWh", f"{frow['Loss (%)']:.2f}%")

        st.download_button(
            "Download Full Report (Excel — expandable detail + flag sections)",
            data=to_excel_bytes(report, consumer_detail, never_billed),
            file_name=f"feeder_wise_loss_{current_month_label}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
else:
    st.info("Upload the installation file, one feeder consumption file, "
            "and two billing report files (current + previous month) to begin.")